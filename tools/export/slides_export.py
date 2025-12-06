"""Slide export utilities."""

from __future__ import annotations

from pathlib import Path
from typing import Dict, List, Optional
from uuid import uuid4

from pptx import Presentation
from pptx.util import Inches, Pt


class SlidesExporter:
    """Creates PPTX decks for presentation-ready summaries."""

    def __init__(self, *, output_dir: str = "/tmp") -> None:
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def export_calendar(self, calendar: Dict[str, object], filename: Optional[str] = None) -> str:
        prs = Presentation()
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        slide.shapes.title.text = f"Content Calendar — {calendar.get('brand', 'Brand')}"
        slide.placeholders[1].text = f"Niche: {calendar.get('niche', 'n/a')} | Days: {calendar.get('days', 'n/a')}"

        bullet_layout = prs.slide_layouts[1]
        ideas: List[Dict[str, object]] = calendar.get("ideas", [])  # type: ignore[assignment]
        for idea in ideas:
            slide = prs.slides.add_slide(bullet_layout)
            slide.shapes.title.text = f"Day {idea.get('day_index')}: {idea.get('theme')}"
            body = slide.shapes.placeholders[1].text_frame
            body.clear()
            body.word_wrap = True
            body.margin_bottom = Inches(0.2)
            p = body.add_paragraph()
            p.text = f"Hook: {idea.get('hook', '')}"
            p.font.size = Pt(14)
            body.add_paragraph().text = f"CTA: {idea.get('call_to_action', '')}"
            body.add_paragraph().text = f"Channel: {idea.get('primary_channel', '')}"
            assets = ', '.join(idea.get('assets_needed', []))
            if assets:
                body.add_paragraph().text = f"Assets: {assets}"
            notes = idea.get('notes')
            if notes:
                body.add_paragraph().text = f"Notes: {notes}"

        output_path = self.output_dir / (filename or f"content-calendar-{uuid4().hex}.pptx")
        prs.save(str(output_path))
        return str(output_path)


if __name__ == "__main__":
    exporter = SlidesExporter()
    sample = {
        "brand": "APAS Fitness",
        "niche": "fitness",
        "days": 3,
        "ideas": [
            {
                "day_index": 1,
                "theme": "Leg Day",
                "hook": "Stop skipping legs",
                "call_to_action": "Join the 30-day challenge",
                "primary_channel": "instagram",
                "assets_needed": ["reel", "caption"],
                "notes": "Highlight testimonials",
            }
        ],
    }
    print(exporter.export_calendar(sample))